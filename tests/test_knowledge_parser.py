"""文档解析测试：验证结构（标题 / 列表 / 表格 / 页码 / 文档顺序）在解析阶段被保留。

纯函数测试，不依赖 MySQL / Qdrant / ARQ。
"""

import io

import pytest
from docx import Document as DocxDocument

import app.modules.knowledge.parser as parser_module
from app.core.exceptions import ValidateError
from app.modules.knowledge.parser import SUPPORTED_EXTS, BlockKind, parse_document


def _kinds(blocks) -> list[str]:
    return [b.kind for b in blocks]


# ---- 纯文本 ----


def test_parse_txt_splits_paragraphs() -> None:
    blocks = parse_document("第一段。\n\n第二段。".encode(), "txt")
    assert [b.text for b in blocks] == ["第一段。", "第二段。"]
    assert _kinds(blocks) == [BlockKind.PARAGRAPH, BlockKind.PARAGRAPH]


def test_parse_txt_normalizes_crlf() -> None:
    blocks = parse_document("第一段。\r\n\r\n第二段。".encode(), "txt")
    assert [b.text for b in blocks] == ["第一段。", "第二段。"]


def test_parse_empty_text_yields_no_blocks() -> None:
    assert parse_document(b"   \n\n  \n", "txt") == []


def test_parse_unsupported_ext() -> None:
    # 文件层白名单比解析层宽（如 xlsx），解析层必须明确拒绝
    with pytest.raises(ValidateError):
        parse_document(b"data", "xlsx")
    with pytest.raises(ValidateError):
        parse_document(b"data", None)


def test_supported_exts_cover_parser_contract() -> None:
    assert {"txt", "md", "pdf", "docx"} == SUPPORTED_EXTS


# ---- Markdown ----


def test_parse_markdown_heading_levels() -> None:
    md = "# 员工手册\n\n## 二、考勤\n\n每天打卡两次。\n\n### 2.1 迟到\n\n迟到需说明原因。"
    blocks = parse_document(md.encode(), "md")
    headings = [(b.text, b.level) for b in blocks if b.kind == BlockKind.HEADING]
    assert headings == [("员工手册", 1), ("二、考勤", 2), ("2.1 迟到", 3)]


def test_parse_markdown_list_kept_as_one_block() -> None:
    md = "## 清单\n\n- 第一项\n- 第二项\n- 第三项\n"
    blocks = parse_document(md.encode(), "md")
    lists = [b for b in blocks if b.kind == BlockKind.LIST]
    assert len(lists) == 1
    assert lists[0].text.splitlines() == ["- 第一项", "- 第二项", "- 第三项"]


def test_parse_markdown_gfm_table() -> None:
    md = "| 姓名 | 部门 |\n| --- | --- |\n| 张三 | 研发 |\n| 李四 | 市场 |\n"
    blocks = parse_document(md.encode(), "md")
    tables = [b for b in blocks if b.kind == BlockKind.TABLE]
    assert len(tables) == 1
    assert tables[0].text.splitlines() == [
        "| 姓名 | 部门 |",
        "| --- | --- |",
        "| 张三 | 研发 |",
        "| 李四 | 市场 |",
    ]


def test_parse_markdown_code_fence_not_misdetected() -> None:
    """代码围栏内的 `#` 与 `|` 不得被误判为标题 / 表格。"""
    md = '## 示例\n\n```python\n# 这不是标题\nprint("| a | b |")\n```\n\n后续段落。'
    blocks = parse_document(md.encode(), "md")
    assert [b.text for b in blocks if b.kind == BlockKind.HEADING] == ["示例"]
    assert [b for b in blocks if b.kind == BlockKind.TABLE] == []
    codes = [b for b in blocks if b.kind == BlockKind.CODE]
    assert len(codes) == 1
    assert "# 这不是标题" in codes[0].text
    assert blocks[-1].text == "后续段落。"


# ---- DOCX ----


def _build_docx() -> bytes:
    doc = DocxDocument()
    doc.add_heading("员工手册", level=1)
    doc.add_paragraph("本手册适用范围为全体员工。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "标准"
    table.cell(1, 0).text = "打卡"
    table.cell(1, 1).text = "每日两次"
    doc.add_paragraph("表格之后的说明段落。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_parse_docx_keeps_document_order() -> None:
    """回归：表格必须留在它原本的位置，不能被排到所有段落之后。"""
    blocks = parse_document(_build_docx(), "docx")
    assert _kinds(blocks) == [
        BlockKind.HEADING,
        BlockKind.PARAGRAPH,
        BlockKind.TABLE,
        BlockKind.PARAGRAPH,
    ]
    assert blocks[2].text.splitlines() == ["项目 | 标准", "打卡 | 每日两次"]
    assert blocks[3].text == "表格之后的说明段落。"


def test_parse_docx_heading_level() -> None:
    blocks = parse_document(_build_docx(), "docx")
    assert blocks[0].level == 1


# ---- PDF（用假 PdfReader，避免引入生成 PDF 的依赖） ----


class _FakePage:
    def __init__(self, text: str) -> None:
        self._text = text

    def extract_text(self) -> str:
        return self._text


class _FakeReader:
    def __init__(self, stream) -> None:
        self.pages = [
            _FakePage("第一页第一段。\n\n第一页第二段。"),
            _FakePage("第二页内容。"),
            _FakePage(""),  # 空页应被跳过
        ]


def test_parse_pdf_keeps_page_number(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(parser_module, "PdfReader", _FakeReader)
    blocks = parse_document(b"ignored", "pdf")
    assert [(b.text, b.page_no) for b in blocks] == [
        ("第一页第一段。", 1),
        ("第一页第二段。", 1),
        ("第二页内容。", 2),
    ]
    assert _kinds(blocks) == [BlockKind.PARAGRAPH] * 3

