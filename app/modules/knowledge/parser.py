"""文档解析与结构感知切分（RAG 索引前置处理）。

原实现是「整篇文本 → 一把字符切分」，标题层级、列表、表格、页码在解析阶段就丢了。
本模块改为两段式：

    文件字节
      → parse_document()  解析为**结构块列表**（heading/paragraph/list/table/code + page_no）
      → split_blocks()    按结构边界分组为 chunk 草稿（带 section_path / page_no / chunk_type）

设计约束：
- 不引入文档解析框架，只用现有 pypdf / python-docx + 正则；
- 保留 chunk_size / chunk_overlap 配置：仅当单个结构块超长时才回退字符切分；
- 结构块内部不切断（标题必开新块、列表整组、表格整体保留到超长为止）。
"""

import io
import re
from dataclasses import dataclass
from enum import StrEnum

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from app.common.enums import ChunkType
from app.core.config import get_settings
from app.core.exceptions import ValidateError

# 支持解析的扩展名（不含点、小写）
SUPPORTED_EXTS = {"txt", "md", "pdf", "docx"}

# 章节路径分隔符
SECTION_SEP = " > "

# 空行（段落边界）
_BLANK_LINE = re.compile(r"\n[ \t]*\n(?:[ \t]*\n)*")

# Markdown 结构
_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_MD_FENCE = re.compile(r"^\s*(```|~~~)")
_MD_LIST = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s+")
_MD_TABLE_ROW = re.compile(r"^\s*\|.*\|\s*$")
_MD_TABLE_SEP = re.compile(r"^\s*\|(?:\s*:?-{2,}:?\s*\|)+\s*$")

# docx 标题样式（兼容中文 Word：Heading 1 / 标题 1）
_DOCX_HEADING = re.compile(r"^(?:Heading|标题)\s*(\d+)$", re.IGNORECASE)

# 超长结构块回退字符切分时的分隔符（优先中文标点与换行，避免把语义切碎）
SPLIT_SEPARATORS = ["\n\n", "\n", "。", "！", "？", "；", "，", " ", ""]


class BlockKind(StrEnum):
    """解析出的结构块类型（内部概念，不落库）。"""

    HEADING = "heading"
    PARAGRAPH = "paragraph"
    LIST = "list"
    TABLE = "table"
    CODE = "code"


@dataclass
class ParsedBlock:
    """解析产出的结构块（切分的最小单元）。"""

    text: str
    kind: str = BlockKind.PARAGRAPH
    level: int | None = None      # 标题层级 1..6（仅 heading）
    page_no: int | None = None    # 来源页码（仅分页文档，如 PDF）


@dataclass
class ChunkDraft:
    """结构感知切分产出的 chunk 草稿（落库前形态）。"""

    content: str
    section_path: str | None = None
    page_no: int | None = None
    chunk_type: str = ChunkType.TEXT.value

    @property
    def embedding_text(self) -> str:
        """用于向量化的文本：章节路径 + 正文。

        切片后丢失的标题上下文在此补回，使「某章节的第 N 个片段」在向量空间里
        仍然带有章节语义；展示用的 `content` 保持干净、不带前缀。
        """
        if self.section_path:
            return f"{self.section_path}\n{self.content}"
        return self.content


# ---- 解析：文件字节 → 结构块列表 ----


def parse_document(data: bytes, ext: str | None) -> list[ParsedBlock]:
    """解析文档字节为结构块列表。ext 不含点、小写；不支持的格式抛 ValidateError。"""
    ext = (ext or "").lower()
    if ext not in SUPPORTED_EXTS:
        raise ValidateError(f"不支持解析的文件类型: {ext or '未知'}")
    if ext == "md":
        return _parse_markdown(_decode_text(data))
    if ext == "txt":
        return _parse_plain(_decode_text(data))
    if ext == "pdf":
        return _parse_pdf(data)
    return _parse_docx(data)


def _decode_text(data: bytes) -> str:
    """字节 → 文本（统一换行符，便于后续按行解析）。"""
    return data.decode("utf-8", errors="replace").replace("\r\n", "\n").replace("\r", "\n")


def _split_paragraphs(text: str) -> list[str]:
    """按空行切分为段落（丢弃空白段）。"""
    return [p.strip() for p in _BLANK_LINE.split(text) if p.strip()]


def _parse_plain(text: str) -> list[ParsedBlock]:
    """纯文本 → 段落块（无标题结构可识别）。"""
    return [ParsedBlock(text=p, kind=BlockKind.PARAGRAPH) for p in _split_paragraphs(text)]


def _is_list_continuation(line: str) -> bool:
    """列表项的缩进续行（非空且以空白开头）。"""
    return bool(line.strip()) and line[:1] in (" ", "\t")


def _parse_markdown(text: str) -> list[ParsedBlock]:
    """Markdown → 结构块（标题 / 段落 / 列表 / 表格 / 代码围栏）。

    代码围栏优先识别并整体保留，避免把代码里的 `#`、`|` 误判为标题或表格。
    """
    blocks: list[ParsedBlock] = []
    lines = text.split("\n")
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            blocks.append(ParsedBlock(text="\n".join(buffer).strip(), kind=BlockKind.PARAGRAPH))
            buffer.clear()

    i = 0
    while i < len(lines):
        line = lines[i]

        # 代码围栏：整体保留，内部不做结构识别
        if _MD_FENCE.match(line):
            flush()
            fence = line.strip()[:3]
            code = [line]
            i += 1
            while i < len(lines):
                code.append(lines[i])
                closed = lines[i].strip().startswith(fence)
                i += 1
                if closed:
                    break
            blocks.append(ParsedBlock(text="\n".join(code).strip(), kind=BlockKind.CODE))
            continue

        heading = _MD_HEADING.match(line)
        if heading:
            flush()
            blocks.append(
                ParsedBlock(
                    text=heading.group(2).strip(),
                    kind=BlockKind.HEADING,
                    level=len(heading.group(1)),
                )
            )
            i += 1
            continue

        # GFM 表格：当前行是 | ... | 且下一行是 |---|---|
        if _MD_TABLE_ROW.match(line) and i + 1 < len(lines) and _MD_TABLE_SEP.match(lines[i + 1]):
            flush()
            rows: list[str] = []
            while i < len(lines) and _MD_TABLE_ROW.match(lines[i]):
                rows.append(lines[i].strip())
                i += 1
            blocks.append(ParsedBlock(text="\n".join(rows), kind=BlockKind.TABLE))
            continue

        # 列表：连续列表行 + 缩进续行合并为一个块（保证列表不被切断）
        if _MD_LIST.match(line):
            flush()
            items: list[str] = []
            while i < len(lines) and (_MD_LIST.match(lines[i]) or _is_list_continuation(lines[i])):
                items.append(lines[i].rstrip())
                i += 1
            blocks.append(ParsedBlock(text="\n".join(items).strip(), kind=BlockKind.LIST))
            continue

        if not line.strip():
            flush()
            i += 1
            continue

        buffer.append(line)
        i += 1

    flush()
    return [b for b in blocks if b.text]


def _parse_pdf(data: bytes) -> list[ParsedBlock]:
    """PDF → 段落块（逐页提取，每块带 page_no）。

    PDF 无可靠的结构信息，故只保留页码，不做标题推断（避免误判污染 section_path）。
    """
    reader = PdfReader(io.BytesIO(data))
    blocks: list[ParsedBlock] = []
    for page_no, page in enumerate(reader.pages, start=1):
        for paragraph in _split_paragraphs(page.extract_text() or ""):
            blocks.append(ParsedBlock(text=paragraph, kind=BlockKind.PARAGRAPH, page_no=page_no))
    return blocks


def _iter_docx_body(document) -> list[Paragraph | Table]:
    """按文档真实顺序遍历 docx 正文（段落与表格交错）。

    不能用 `document.paragraphs` + `document.tables` 两次遍历：那会把所有表格
    排到所有段落之后，破坏表格与上下文的对应关系。
    """
    items: list[Paragraph | Table] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            items.append(Paragraph(child, document))
        elif child.tag == qn("w:tbl"):
            items.append(Table(child, document))
    return items


def _docx_heading_level(paragraph: Paragraph) -> int | None:
    """从段落样式推断标题层级（Heading 1 / 标题 1）；非标题返回 None。"""
    style = paragraph.style
    if style is None or not style.name:
        return None
    match = _DOCX_HEADING.match(style.name.strip())
    if match is None:
        return None
    return min(int(match.group(1)), 6)


def _parse_docx(data: bytes) -> list[ParsedBlock]:
    """docx → 结构块（标题 / 段落 / 表格，保持文档顺序）。"""
    document = Document(io.BytesIO(data))
    blocks: list[ParsedBlock] = []
    for item in _iter_docx_body(document):
        if isinstance(item, Table):
            rows = [
                " | ".join(cell.text.strip() for cell in row.cells).strip() for row in item.rows
            ]
            text = "\n".join(r for r in rows if r.strip(" |"))
            if text:
                blocks.append(ParsedBlock(text=text, kind=BlockKind.TABLE))
            continue
        text = item.text.strip()
        if not text:
            continue
        level = _docx_heading_level(item)
        if level is None:
            blocks.append(ParsedBlock(text=text, kind=BlockKind.PARAGRAPH))
        else:
            blocks.append(ParsedBlock(text=text, kind=BlockKind.HEADING, level=level))
    return blocks


# ---- 切分：结构块 → chunk 草稿 ----


def _push_section(section: list[tuple[int, str]], block: ParsedBlock) -> None:
    """维护标题栈：弹出层级不低于当前标题的旧标题。"""
    level = block.level or 1
    while section and section[-1][0] >= level:
        section.pop()
    section.append((level, block.text))


def _section_path(section: list[tuple[int, str]]) -> str | None:
    """标题栈 → 章节路径（如「员工手册 > 二、考勤」）；无标题时返回 None。"""
    if not section:
        return None
    return SECTION_SEP.join(text for _, text in section)


def _heading_only(buffer: list[ParsedBlock]) -> bool:
    """缓冲里是否只有标题（此时不应落盘，否则产出纯标题碎块）。"""
    return len(buffer) == 1 and buffer[0].kind == BlockKind.HEADING


def _chunk_type(blocks: list[ParsedBlock]) -> str:
    """由块类型集合推断 chunk 类型。"""
    kinds = {b.kind for b in blocks}
    if kinds == {BlockKind.TABLE}:
        return ChunkType.TABLE.value
    if BlockKind.TABLE in kinds:
        return ChunkType.MIXED.value
    return ChunkType.TEXT.value


def _table_header_size(lines: list[str]) -> int:
    """表格头部行数：GFM 表格为「表头 + 分隔行」（2 行），其余表格按 1 行表头处理。"""
    if len(lines) >= 2 and _MD_TABLE_SEP.match(lines[1]):
        return 2
    return 1


def _split_oversized_table(content: str, chunk_size: int) -> list[str]:
    """超大表格按整行切分，并在每个子表前重复表头（表头 + 分隔行）。

    表格的行是有语义的原子单位：切在行中间会破坏数据，切掉表头会让子表失去列含义
    （例如只有「| 迟到豁免 | 每月 3 次 |」，读不出这三列分别是什么）。
    此路径不套用 chunk_overlap —— 保证行完整与子表自解释，比字符重叠更有价值。
    """
    lines = [line for line in content.splitlines() if line.strip()]
    header_size = _table_header_size(lines)
    if len(lines) <= header_size:
        return [content]
    header = lines[:header_size]
    header_text = "\n".join(header)
    pieces: list[str] = []
    current: list[str] = []
    current_len = len(header_text)
    for row in lines[header_size:]:
        if current and current_len + len(row) + 1 > chunk_size:
            pieces.append("\n".join(header + current))
            current = []
            current_len = len(header_text)
        current.append(row)
        current_len += len(row) + 1
    if current:
        pieces.append("\n".join(header + current))
    return pieces


def _make_table_drafts(
    blocks: list[ParsedBlock],
    section_path: str | None,
    page_no: int | None,
    chunk_type: str,
    chunk_size: int,
) -> list[ChunkDraft]:
    """超长「含表格」缓冲的切分：表格按行切分并重复表头。

    非表格前缀（标题 / 引导段落）并入**第一片**，避免产出「只有标题」的碎块；
    第 2 片起是纯表格，故类型标为 table（首片保留原类型，可能是 mixed）。
    """
    prefix = "\n\n".join(b.text for b in blocks if b.kind != BlockKind.TABLE).strip()
    table_text = "\n".join(b.text for b in blocks if b.kind == BlockKind.TABLE)
    pieces = _split_oversized_table(table_text, chunk_size)
    if prefix:
        pieces = [f"{prefix}\n\n{pieces[0]}", *pieces[1:]] if pieces else [prefix]
    return [
        ChunkDraft(
            content=piece,
            section_path=section_path,
            page_no=page_no,
            chunk_type=chunk_type if index == 0 else ChunkType.TABLE.value,
        )
        for index, piece in enumerate(pieces)
        if piece.strip()
    ]


def _make_drafts(
    blocks: list[ParsedBlock], section_path: str | None, chunk_size: int, overlap: int
) -> list[ChunkDraft]:
    """把一个缓冲的结构块合成 chunk；超长时按块类型选择切分策略。"""
    # 块之间用空行连接：保留段落边界（段落边界本身也是文档结构）
    content = "\n\n".join(b.text for b in blocks).strip()
    if not content:
        return []
    # 跨页块取起始页（PDF 已在 split_blocks 里按页断开，多页仅可能出现在非分页文档）
    page_no = next((b.page_no for b in blocks if b.page_no is not None), None)
    chunk_type = _chunk_type(blocks)
    if len(content) <= chunk_size:
        return [
            ChunkDraft(
                content=content,
                section_path=section_path,
                page_no=page_no,
                chunk_type=chunk_type,
            )
        ]
    # 含表格：表格按行切分 + 重复表头（表格行是语义原子，表头决定列含义）
    if any(b.kind == BlockKind.TABLE for b in blocks):
        return _make_table_drafts(blocks, section_path, page_no, chunk_type, chunk_size)
    # 其他超长结构块（长段落 / 列表）：回退字符切分，保留 chunk_size / chunk_overlap 配置
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=overlap,
        separators=SPLIT_SEPARATORS,
    )
    return [
        ChunkDraft(
            content=piece.strip(),
            section_path=section_path,
            page_no=page_no,
            chunk_type=chunk_type,
        )
        for piece in splitter.split_text(content)
        if piece.strip()
    ]


def split_blocks(blocks: list[ParsedBlock]) -> list[ChunkDraft]:
    """按结构边界把结构块分组成 chunk 草稿（chunk_size / chunk_overlap 仍生效）。

    分组规则：
    - 标题必开新块，使章节边界与块边界对齐（标题与其内容同块）；
    - 换页断开（PDF），保证 page_no 精确；
    - 累计长度超过 chunk_size 时先落盘；但缓冲里只有标题时不落盘
      （否则会产出「只有标题」的碎块）；
    - 纯标题缓冲直接丢弃 —— 标题的价值已由 section_path 承载，单独成块只是噪声。
    """
    settings = get_settings()
    chunk_size = max(1, settings.chunk_size)
    overlap = max(0, settings.chunk_overlap)

    drafts: list[ChunkDraft] = []
    section: list[tuple[int, str]] = []
    buffer: list[ParsedBlock] = []
    buffer_len = 0
    buffer_page: int | None = None

    def flush() -> None:
        nonlocal buffer, buffer_len, buffer_page
        if buffer and not _heading_only(buffer):
            drafts.extend(_make_drafts(buffer, _section_path(section), chunk_size, overlap))
        buffer = []
        buffer_len = 0
        buffer_page = None

    for block in blocks:
        if block.kind == BlockKind.HEADING:
            flush()
            _push_section(section, block)
            buffer = [block]
            buffer_len = len(block.text)
            buffer_page = block.page_no
            continue
        # 换页：按页断开，保证 page_no 精确
        if block.page_no is not None and buffer_page is not None and block.page_no != buffer_page:
            flush()
        # 超长：先落盘（缓冲里只有标题时不落盘，让标题跟内容走）
        if buffer and buffer_len + len(block.text) > chunk_size and not _heading_only(buffer):
            flush()
        if not buffer:
            buffer_page = block.page_no
        buffer.append(block)
        buffer_len += len(block.text)

    flush()
    return drafts

